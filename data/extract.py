"""Extract index data from index.txt and generate an xlsx and a json file."""

import json
from pathlib import Path

from openpyxl import Workbook

BASE_DIR = Path(r"C:\Users\EDY\Desktop\26年第三季度\资讯关联基金产品\code")
SRC = BASE_DIR / "data" / "index.txt"
XLSX_OUT = BASE_DIR / "data" / "index.xlsx"
JSON_OUT = BASE_DIR / "data" / "index.json"


def main() -> None:
    raw = SRC.read_text(encoding="utf-8")
    payload = json.loads(raw)
    records = payload["json"][0]["data"]

    # Build a flat list with consistent fields
    flat = [
        {
            "INDEX_CODE": r.get("INDEX_CODE") or "",
            "INDEX_NAME": r.get("INDEX_NAME") or "",
            "INDEX_TYPE_NAME": r.get("INDEX_TYPE_NAME") or "",
            "ETF_COUNT": r.get("ETF_COUNT") or 0,
            "HTF_ETF_COUNT": r.get("HTF_ETF_COUNT") or 0,
            "INDEX_ABSTRACT": r.get("INDEX_ABSTRACT") or "",
            "INDEX_REMARK": r.get("INDEX_REMARK") or "",
        }
        for r in records
    ]
    print(f"Extracted {len(flat)} records")

    # ---- Write JSON ----
    JSON_OUT.write_text(
        json.dumps(flat, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(f"JSON written to: {JSON_OUT}")

    # ---- Write XLSX ----
    wb = Workbook()
    ws = wb.active
    ws.title = "Indexes"
    headers = [
        "INDEX_CODE", "INDEX_NAME", "INDEX_TYPE_NAME",
        "ETF_COUNT", "HTF_ETF_COUNT", "INDEX_ABSTRACT", "INDEX_REMARK",
    ]
    ws.append(headers)
    for row in flat:
        ws.append([row[h] for h in headers])

    # Adjust column widths so the content is readable
    widths = {"A": 14, "B": 24, "C": 14, "D": 10, "E": 14, "F": 60, "G": 80}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    # Wrap text and align top for long fields
    for cell in ws["F"][1:] + ws["G"][1:]:
        cell.alignment = cell.alignment.copy(wrap_text=True, vertical="top")

    wb.save(XLSX_OUT)
    print(f"XLSX written to: {XLSX_OUT}")


if __name__ == "__main__":
    main()